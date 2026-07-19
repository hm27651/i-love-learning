import importlib
import io
import json
import os
import shutil
import sqlite3
import tempfile
import time
import unittest
import re
from pathlib import Path


class StudyAppTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.original_study_data_dir = os.environ.get("STUDY_DATA_DIR")
        cls.original_h3cse_data_dir = os.environ.get("H3CSE_DATA_DIR")
        cls.original_study_backup_dir = os.environ.get("STUDY_BACKUP_DIR")
        os.environ["STUDY_DATA_DIR"] = cls.temp.name
        os.environ["H3CSE_DATA_DIR"] = cls.temp.name
        os.environ["STUDY_BACKUP_DIR"] = str(Path(cls.temp.name) / "backups")
        cls.mod = importlib.import_module("app")
        if cls.mod.DATA_DIR.resolve() != Path(cls.temp.name).resolve():
            raise RuntimeError(f"测试数据目录隔离失败：{cls.mod.DATA_DIR}")
        cls.mod.app.config.update(TESTING=True)
        cls.mod.BACKUP_DIR = Path(cls.temp.name) / "backups"

    @classmethod
    def tearDownClass(cls):
        from app_runtime import close_file_logging

        close_file_logging(cls.mod.app.logger, Path(cls.temp.name))
        cls.mod.background_executor.shutdown(wait=True)
        cls.temp.cleanup()
        for name, value in (
            ("STUDY_DATA_DIR", cls.original_study_data_dir),
            ("H3CSE_DATA_DIR", cls.original_h3cse_data_dir),
            ("STUDY_BACKUP_DIR", cls.original_study_backup_dir),
        ):
            if value is None:
                os.environ.pop(name, None)
            else:
                os.environ[name] = value

    def setUp(self):
        self.client = self.mod.app.test_client()
        with self.mod.db() as conn:
            self.project_id = conn.execute("SELECT id FROM learning_projects ORDER BY id LIMIT 1").fetchone()["id"]
            conn.execute("UPDATE learning_projects SET project_type='practical_certification',practice_alias='HCL实验' WHERE id=?", (self.project_id,))
            conn.execute("UPDATE project_modules SET enabled=1 WHERE project_id=?", (self.project_id,))
            for key, value in self.mod.PROJECT_SETTING_DEFAULTS.items():
                conn.execute("""INSERT INTO project_settings(project_id,key,value) VALUES (?,?,?)
                  ON CONFLICT(project_id,key) DO UPDATE SET value=excluded.value""", (self.project_id, key, value))
            conn.execute("""INSERT INTO project_settings(project_id,key,value) VALUES (?,'onboarding_completed','1')
              ON CONFLICT(project_id,key) DO UPDATE SET value='1'""", (self.project_id,))
            conn.execute("DELETE FROM attempts")
            conn.execute("DELETE FROM question_progress")
            conn.execute("DELETE FROM practice_sessions")
            conn.execute("DELETE FROM mock_exams")
            conn.execute("DELETE FROM import_candidates")
            conn.execute("DELETE FROM questions")
            conn.execute("DELETE FROM labs")
            conn.execute("DELETE FROM import_jobs")
            conn.execute("DELETE FROM export_jobs")
            conn.execute("DELETE FROM source_documents")
            conn.execute("DELETE FROM knowledge_points")
            conn.execute("DELETE FROM chapters")
            conn.execute("DELETE FROM subjects")
            subject = conn.execute("INSERT INTO subjects(project_id,name,code) VALUES (?,?,?)",
                                   (self.project_id, "测试科目", "TEST")).lastrowid
            chapter = conn.execute("INSERT INTO chapters(subject_id,name,is_core) VALUES (?,?,1)",
                                   (subject, "测试章节")).lastrowid
            conn.execute("INSERT INTO knowledge_points(chapter_id,name) VALUES (?,?)", (chapter, "测试知识点"))

    def wait_for_job(self, job_id, statuses, timeout=8):
        deadline = time.time() + timeout
        while time.time() < deadline:
            with self.mod.db() as conn:
                row = conn.execute("SELECT * FROM import_jobs WHERE id=?", (job_id,)).fetchone()
                if row and row["status"] in statuses:
                    return row
            time.sleep(0.05)
        self.fail(f"job {job_id} did not reach {statuses}")

    def wait_for_export(self, job_id, statuses, timeout=8):
        deadline = time.time() + timeout
        while time.time() < deadline:
            with self.mod.db() as conn:
                row = conn.execute("SELECT * FROM export_jobs WHERE id=?", (job_id,)).fetchone()
                if row and row["status"] in statuses:
                    return row
            time.sleep(0.05)
        self.fail(f"export {job_id} did not reach {statuses}")

    def create_export(self, scope_type="project", scope_id=None, include_drafts=False):
        with self.mod.db() as conn:
            before = {row["id"] for row in conn.execute("SELECT id FROM export_jobs")}
        data = {"scope_type": scope_type}
        if scope_id is not None:
            data["scope_id"] = str(scope_id)
        if include_drafts:
            data["include_drafts"] = "on"
        response = self.client.post("/exports", data=data)
        self.assertEqual(response.status_code, 302)
        with self.mod.db() as conn:
            created = [row["id"] for row in conn.execute("SELECT id FROM export_jobs") if row["id"] not in before]
            self.assertEqual(len(created), 1)
            job_id = created[0]
        return job_id, self.wait_for_export(job_id, {"completed", "failed"})

    def upload_csv_job(self, stem="通用导入测试题"):
        csv_data = f"题型,题干,选项A,选项B,选项C,选项D,选项E,选项F,选项G,选项H,答案,解析,章节,知识点,原题号\n单选,{stem},正确项,错误项,,,,,,,A,测试解析,导入章节,导入知识点,CSV-1\n"
        response = self.client.post("/imports", data={"file": (io.BytesIO(csv_data.encode("utf-8-sig")), "bank.csv")}, content_type="multipart/form-data")
        job_id = response.headers["Location"].rstrip("/").split("/")[-1]
        self.wait_for_job(job_id, {"waiting_target"})
        with self.mod.db() as conn:
            subject_id = conn.execute("SELECT id FROM subjects WHERE project_id=? ORDER BY id LIMIT 1", (self.project_id,)).fetchone()["id"]
        self.client.post(f"/imports/{job_id}/target", data={"project_id": self.project_id, "subject_id": subject_id})
        return job_id, self.wait_for_job(job_id, {"ready", "blocked", "failed"})

    def add_question(self, qtype="single", answer=None, status="verified"):
        with self.mod.db() as conn:
            point = conn.execute("SELECT id FROM knowledge_points LIMIT 1").fetchone()["id"]
            answers = answer or (["A", "C"] if qtype == "multiple" else (["参考答案"] if qtype in {"fill", "short"} else ["A"]))
            options = ["选项一", "选项二", "选项三"] if qtype in {"single", "multiple"} else []
            cur = conn.execute("""INSERT INTO questions(knowledge_point_id,type,stem,options_json,answer_json,explanation,difficulty,status,created_at,updated_at)
                VALUES (?,?,?,?,?,?,?,?,?,?)""", (point, qtype, f"{qtype} 测试题", json.dumps(options, ensure_ascii=False),
                json.dumps(answers, ensure_ascii=False), "测试解析", 2, status, self.mod.now_iso(), self.mod.now_iso()))
            return cur.lastrowid

    def control_token(self, table, item_id):
        with self.mod.db() as conn:
            return conn.execute(f"SELECT control_token FROM {table} WHERE id=?", (item_id,)).fetchone()["control_token"]

    def create_tree(self, subject_name, chapter_name=None, point_name=None, *, project_id=None, core=0):
        project_id = project_id or self.project_id
        with self.mod.db() as conn:
            subject = conn.execute("INSERT INTO subjects(project_id,name,code) VALUES (?,?,?)",
                                   (project_id, subject_name, "")).lastrowid
            chapter = point = None
            if chapter_name is not None:
                chapter = conn.execute("INSERT INTO chapters(subject_id,name,is_core) VALUES (?,?,?)",
                                       (subject, chapter_name, core)).lastrowid
            if point_name is not None:
                point = conn.execute("INSERT INTO knowledge_points(chapter_id,name) VALUES (?,?)",
                                     (chapter, point_name)).lastrowid
            return subject, chapter, point

    def test_all_main_pages_render(self):
        self.add_question()
        paths = ["/", "/progress", "/knowledge", "/questions", "/questions/new", "/practice", "/review", "/labs", "/labs/new", "/plans", "/mock", "/imports", "/settings", "/data-management", "/health"]
        for path in paths:
            response = self.client.get(path)
            self.assertEqual(response.status_code, 200, path)

    def test_sidebar_keeps_project_switcher_without_collapsed_initial_badge(self):
        page = self.client.get("/").get_data(as_text=True)
        self.assertIn('id="sidebar-project"', page)
        self.assertNotIn("collapsed-project-pill", page)

    def test_sqlite_runtime_and_background_queue_are_shared(self):
        with self.mod.db() as conn:
            self.assertEqual(conn.execute("PRAGMA journal_mode").fetchone()[0], "wal")
            self.assertEqual(conn.execute("PRAGMA busy_timeout").fetchone()[0], 30000)
        self.assertIs(self.mod.import_queue.executor, self.mod.export_queue.executor)

    def test_create_app_builds_an_isolated_runtime_without_route_back_imports(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            factory_app = self.mod.create_app(
                {
                    "TESTING": True,
                    "SECRET_KEY": "factory-test",
                    "STUDY_DATA_DIR": root / "data",
                    "STUDY_DB_PATH": root / "data" / "factory.db",
                    "STUDY_BACKUP_DIR": root / "backups",
                }
            )
            try:
                self.assertIsNot(factory_app, self.mod.app)
                self.assertEqual(Path(factory_app.config["STUDY_DATA_DIR"]), (root / "data").resolve())
                with factory_app.app_context():
                    from services.core.runtime_service import db as runtime_db

                    with runtime_db() as conn:
                        self.assertEqual(conn.execute("SELECT MAX(version) FROM schema_migrations").fetchone()[0], 4)
                        self.assertEqual(conn.execute("SELECT COUNT(*) FROM questions").fetchone()[0], 0)
                for route_file in (Path(__file__).resolve().parents[1] / "routes").glob("*.py"):
                    source = route_file.read_text(encoding="utf-8")
                    self.assertNotIn("from app import", source, route_file.name)
                    self.assertNotIn("import app as", source, route_file.name)
            finally:
                factory_app.extensions["study_background_executor"].shutdown(wait=True)
                from app_runtime import close_file_logging

                close_file_logging(factory_app.logger, root / "data")

    def test_first_run_wizard_configures_project_modules_without_affecting_existing_data(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            factory_app = self.mod.create_app(
                {
                    "TESTING": True,
                    "ONBOARDING_IN_TESTS": True,
                    "SECRET_KEY": "onboarding-test",
                    "STUDY_DATA_DIR": root / "data",
                    "STUDY_DB_PATH": root / "data" / "onboarding.db",
                    "STUDY_BACKUP_DIR": root / "backups",
                }
            )
            try:
                client = factory_app.test_client()
                redirect_response = client.get("/")
                self.assertEqual(redirect_response.status_code, 302)
                self.assertTrue(redirect_response.headers["Location"].endswith("/welcome"))
                page = client.get("/welcome")
                self.assertEqual(page.status_code, 200)
                self.assertIn("选择项目模板", page.get_data(as_text=True))

                saved = client.post(
                    "/welcome",
                    data={
                        "name": "网络认证",
                        "project_type": "practical_certification",
                        "start_date": "2026-07-19",
                        "duration_weeks": "12",
                        "practice_alias": "HCL实验",
                    },
                )
                self.assertEqual(saved.status_code, 302)
                with factory_app.app_context():
                    from services.core.runtime_service import db as runtime_db

                    with runtime_db() as conn:
                        project = conn.execute("SELECT * FROM learning_projects").fetchone()
                        modules = dict(conn.execute("SELECT module_key,enabled FROM project_modules"))
                        completed = conn.execute(
                            "SELECT value FROM project_settings WHERE key='onboarding_completed'"
                        ).fetchone()["value"]
                        self.assertEqual(project["name"], "网络认证")
                        self.assertEqual(project["practice_alias"], "HCL实验")
                        self.assertEqual(set(modules.values()), {1})
                        self.assertEqual(completed, "1")
                        self.assertEqual(conn.execute("SELECT COUNT(*) FROM questions").fetchone()[0], 0)
                self.assertEqual(client.get("/mock").status_code, 200)
            finally:
                factory_app.extensions["study_background_executor"].shutdown(wait=True)
                from app_runtime import close_file_logging

                close_file_logging(factory_app.logger, root / "data")

    def test_health_and_security_headers(self):
        response = self.client.get("/health")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["database"], "ok")
        self.assertEqual(response.get_json()["version"], "2.0")
        self.assertEqual(response.headers["X-Content-Type-Options"], "nosniff")
        self.assertEqual(response.headers["X-Frame-Options"], "SAMEORIGIN")

    def test_csrf_is_required_outside_testing_mode(self):
        self.mod.app.config["TESTING"] = False
        try:
            client = self.mod.app.test_client()
            page = client.get("/settings")
            token = re.search(rb'<meta name="csrf-token" content="([^"]+)">', page.data).group(1).decode()
            rejected = client.post("/projects/switch", data={"project_id": self.project_id})
            self.assertEqual(rejected.status_code, 400)
            accepted = client.post("/projects/switch", data={"project_id": self.project_id, "_csrf_token": token})
            self.assertEqual(accepted.status_code, 302)
        finally:
            self.mod.app.config["TESTING"] = True

    def test_learning_data_cleanup_rebuilds_progress_and_creates_backup(self):
        question_id = self.add_question()
        now = self.mod.now_iso()
        with self.mod.db() as conn:
            conn.execute("""INSERT INTO practice_sessions(id,project_id,mode,question_ids_json,current_index,
              started_at,completed_at,status,updated_at) VALUES ('keep',?,'practice',?,1,?,?,'completed',?)""",
              (self.project_id, json.dumps([question_id]), now, now, now))
            question = self.mod.get_question(conn, question_id, self.project_id)
            self.mod.record_attempt(conn, question, "practice", is_correct=True, session_id="keep")
            conn.execute("""INSERT INTO practice_sessions(id,project_id,mode,question_ids_json,current_index,
              started_at,terminated_at,status,updated_at) VALUES ('remove',?,'practice',?,1,?,?,'terminated',?)""",
              (self.project_id, json.dumps([question_id]), now, now, now))
            self.mod.record_attempt(conn, question, "practice", is_correct=False, session_id="remove")
        page = self.client.get("/data-management")
        self.assertIn(b"remove", page.data)
        result = self.client.post("/data-management/cleanup", data={
            "session_key": "practice:remove", "confirmation": "清理所选记录"
        }, follow_redirects=True)
        self.assertEqual(result.status_code, 200)
        with self.mod.db() as conn:
            self.assertIsNone(conn.execute("SELECT 1 FROM practice_sessions WHERE id='remove'").fetchone())
            self.assertIsNotNone(conn.execute("SELECT 1 FROM practice_sessions WHERE id='keep'").fetchone())
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM attempts").fetchone()[0], 1)
            progress = conn.execute("SELECT * FROM question_progress WHERE question_id=?", (question_id,)).fetchone()
            self.assertEqual((progress["mastery_level"], progress["attempts"], progress["correct_attempts"], progress["error_count"]), (1, 1, 1, 0))
        self.assertTrue(any(self.mod.BACKUP_DIR.glob("learning_cleanup_*/data/h3cse.db")))

    def test_manual_backup_restores_to_an_independent_directory(self):
        question_id = self.add_question()
        response = self.client.post("/data-management/backup", follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        snapshot = max(self.mod.BACKUP_DIR.glob("manual_*/data"), key=lambda path: path.stat().st_mtime)
        restore = Path(self.temp.name) / "restore-check"
        if restore.exists():
            shutil.rmtree(restore)
        shutil.copytree(snapshot, restore)
        conn = sqlite3.connect(restore / "h3cse.db")
        try:
            self.assertEqual(conn.execute("PRAGMA integrity_check").fetchone()[0], "ok")
            self.assertEqual(conn.execute("SELECT MAX(version) FROM schema_migrations").fetchone()[0], 4)
            self.assertEqual(conn.execute("SELECT id FROM questions").fetchone()[0], question_id)
        finally:
            conn.close()
        self.assertEqual(
            (Path(self.temp.name) / ".secret_key").read_text(encoding="ascii"),
            (restore / ".secret_key").read_text(encoding="ascii"),
        )
        shutil.rmtree(restore)

    def test_question_creation_supports_multiple_choice(self):
        with self.mod.db() as conn:
            point = conn.execute("SELECT id FROM knowledge_points LIMIT 1").fetchone()["id"]
        response = self.client.post("/questions/new", data={
            "knowledge_point_id": point, "type": "multiple", "stem": "选择正确项",
            "options": "第一项\n第二项\n第三项", "answer": "C,A", "explanation": "解析",
            "difficulty": 3, "source": "官方资料", "version_note": "Comware 7", "status": "verified",
        })
        self.assertEqual(response.status_code, 302)
        with self.mod.db() as conn:
            row = conn.execute("SELECT * FROM questions").fetchone()
            self.assertEqual(json.loads(row["answer_json"]), ["A", "C"])
            self.assertEqual(row["status"], "verified")

    def test_question_list_is_paginated(self):
        for _ in range(55):
            self.add_question()
        first = self.client.get("/questions")
        second = self.client.get("/questions?page=2")
        self.assertEqual(first.status_code, 200)
        self.assertEqual(second.status_code, 200)
        self.assertIn("第 1 / 2 页".encode("utf-8"), first.data)
        self.assertIn("第 2 / 2 页".encode("utf-8"), second.data)
        self.assertEqual(first.data.count(b'name="page"'), 2)
        self.assertEqual(first.data.count("跳转".encode("utf-8")), 2)

    def test_question_bulk_status_and_review_workspace(self):
        first = self.add_question(status="draft")
        second = self.add_question(status="draft")
        review = self.client.get(f"/questions/review/{first}?status=draft&page=1")
        self.assertEqual(review.status_code, 200)
        self.assertIn("审核题目".encode("utf-8"), review.data)
        response = self.client.post("/questions/bulk", data={
            "action": "verified", "question_id": [str(first), str(second)],
            "return_to": "/questions?status=draft&page=2",
        })
        self.assertEqual(response.status_code, 302)
        self.assertTrue(response.headers["Location"].endswith("/questions?status=draft&page=2"))
        with self.mod.db() as conn:
            statuses = {row["status"] for row in conn.execute("SELECT status FROM questions WHERE id IN (?,?)", (first, second))}
            self.assertEqual(statuses, {"verified"})

    def test_question_bulk_can_move_within_project_only(self):
        question_id = self.add_question(status="draft")
        with self.mod.db() as conn:
            chapter = conn.execute("""SELECT c.id FROM chapters c JOIN subjects s ON s.id=c.subject_id
              WHERE s.project_id=? ORDER BY c.id LIMIT 1""", (self.project_id,)).fetchone()["id"]
            target = conn.execute("INSERT INTO knowledge_points(chapter_id,name) VALUES (?,?)", (chapter, "批量移动目标")).lastrowid
        response = self.client.post("/questions/bulk", data={"action": "move", "question_id": [str(question_id)], "target_point_id": target})
        self.assertEqual(response.status_code, 302)
        with self.mod.db() as conn:
            self.assertEqual(conn.execute("SELECT knowledge_point_id FROM questions WHERE id=?", (question_id,)).fetchone()["knowledge_point_id"], target)

    def test_objective_practice_updates_mastery_and_due_date(self):
        qid = self.add_question("single", ["A"])
        response = self.client.post("/practice", data={"count": 1})
        location = response.headers["Location"]
        self.assertEqual(self.client.get(location).status_code, 200)
        session_id = location.rstrip("/").split("/")[-1]
        answer = self.client.post(location, data={"answer": "A", "control_token": self.control_token("practice_sessions", session_id)})
        self.assertIn("回答正确".encode("utf-8"), answer.data)
        with self.mod.db() as conn:
            progress = conn.execute("SELECT * FROM question_progress WHERE question_id=?", (qid,)).fetchone()
            self.assertEqual(progress["mastery_level"], 1)
            self.assertEqual(progress["correct_attempts"], 1)

    def test_practice_pause_preserves_unsubmitted_answer_and_can_resume(self):
        self.add_question("single", ["A"])
        response = self.client.post("/practice", data={"count": 1})
        location = response.headers["Location"]
        session_id = location.rstrip("/").split("/")[-1]
        self.client.get(location)
        token = self.control_token("practice_sessions", session_id)
        paused = self.client.post(f"/practice/{session_id}/pause", json={
            "control_token": token, "selected": ["A"], "draft": "暂存说明",
        })
        self.assertEqual(paused.status_code, 200)
        with self.mod.db() as conn:
            item = conn.execute("SELECT * FROM practice_sessions WHERE id=?", (session_id,)).fetchone()
            self.assertEqual(item["status"], "paused")
            self.assertEqual(json.loads(item["state_json"])["selected"], ["A"])
        page = self.client.get(location)
        self.assertIn("练习已暂停".encode(), page.data)
        self.assertIn(b'value="A" checked', page.data)
        resumed = self.client.post(f"/practice/{session_id}/resume", data={
            "control_token": self.control_token("practice_sessions", session_id),
        })
        self.assertEqual(resumed.status_code, 302)
        with self.mod.db() as conn:
            self.assertEqual(conn.execute("SELECT status FROM practice_sessions WHERE id=?", (session_id,)).fetchone()["status"], "active")

    def test_practice_pause_on_feedback_restores_feedback_without_duplicate_attempt(self):
        self.add_question("single", ["A"])
        location = self.client.post("/practice", data={"count": 1}).headers["Location"]
        session_id = location.rstrip("/").split("/")[-1]
        self.client.get(location)
        token = self.control_token("practice_sessions", session_id)
        answered = self.client.post(location, data={"answer": "A", "control_token": token})
        self.assertIn("回答正确".encode(), answered.data)
        paused = self.client.post(f"/practice/{session_id}/pause", json={"control_token": token})
        self.assertEqual(paused.status_code, 200)
        restored = self.client.get(location)
        self.assertIn("回答正确".encode(), restored.data)
        self.assertIn("练习已暂停".encode(), restored.data)
        with self.mod.db() as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) n FROM attempts WHERE session_id=?", (session_id,)).fetchone()["n"], 1)

    def test_practice_terminate_keeps_completed_attempts_and_frees_mode(self):
        self.add_question("single", ["A"])
        location = self.client.post("/practice", data={"count": 1}).headers["Location"]
        session_id = location.rstrip("/").split("/")[-1]
        self.client.get(location)
        token = self.control_token("practice_sessions", session_id)
        self.client.post(location, data={"answer": "A", "control_token": token})
        result = self.client.post(f"/practice/{session_id}/terminate", data={"control_token": token}, follow_redirects=True)
        self.assertIn("会话已终止".encode(), result.data)
        with self.mod.db() as conn:
            self.assertEqual(conn.execute("SELECT status FROM practice_sessions WHERE id=?", (session_id,)).fetchone()["status"], "terminated")
            self.assertEqual(conn.execute("SELECT COUNT(*) n FROM attempts WHERE session_id=?", (session_id,)).fetchone()["n"], 1)
        self.assertEqual(self.client.post("/practice", data={"count": 1}).status_code, 302)

    def test_only_one_active_session_per_practice_mode(self):
        self.add_question("single", ["A"])
        first = self.client.post("/practice", data={"count": 1})
        self.assertIn("/practice/", first.headers["Location"])
        second = self.client.post("/practice", data={"count": 1}, follow_redirects=True)
        self.assertIn("请先继续或终止旧会话".encode(), second.data)
        with self.mod.db() as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) n FROM practice_sessions WHERE mode='practice' AND status IN ('active','paused')").fetchone()["n"], 1)

    def test_practice_match_count_respects_filters_and_status(self):
        self.add_question("single", status="verified")
        self.add_question("multiple", status="verified")
        self.add_question("single", status="draft")
        self.add_question("single", status="archived")
        with self.mod.db() as conn:
            chapter_id = conn.execute("SELECT id FROM chapters LIMIT 1").fetchone()["id"]
        self.assertEqual(self.client.get("/api/practice/count").get_json()["count"], 2)
        self.assertEqual(self.client.get("/api/practice/count?type=single").get_json()["count"], 1)
        self.assertEqual(self.client.get(f"/api/practice/count?chapter_id={chapter_id}").get_json()["count"], 2)
        self.assertEqual(self.client.get("/api/practice/count?chapter_id=999999").get_json()["count"], 0)

    def test_all_practice_mode_includes_every_match_and_count_mode_stays_capped(self):
        for _ in range(105):
            self.add_question("single", status="verified")
        response = self.client.post("/practice", data={"selection_mode": "all", "count": 10})
        self.assertEqual(response.status_code, 302)
        with self.mod.db() as conn:
            session = conn.execute("SELECT * FROM practice_sessions ORDER BY started_at DESC LIMIT 1").fetchone()
            ids = json.loads(session["question_ids_json"])
            self.assertEqual(len(ids), 105)
            self.assertEqual(len(set(ids)), 105)
            conn.execute("DELETE FROM practice_sessions")
        self.client.post("/practice", data={"selection_mode": "count", "count": 500})
        with self.mod.db() as conn:
            session = conn.execute("SELECT * FROM practice_sessions LIMIT 1").fetchone()
            self.assertEqual(len(json.loads(session["question_ids_json"])), 100)

    def test_all_practice_mode_does_not_create_empty_session(self):
        response = self.client.post("/practice", data={"selection_mode": "all"})
        self.assertEqual(response.status_code, 302)
        self.assertTrue(response.headers["Location"].endswith("/practice"))
        with self.mod.db() as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) n FROM practice_sessions").fetchone()["n"], 0)

    def test_wrong_answer_resets_mastery_and_counts_error(self):
        qid = self.add_question("single", ["A"])
        with self.mod.db() as conn:
            question = self.mod.get_question(conn, qid)
            self.mod.record_attempt(conn, question, "practice", is_correct=True)
            self.mod.record_attempt(conn, question, "practice", is_correct=False)
            progress = conn.execute("SELECT * FROM question_progress WHERE question_id=?", (qid,)).fetchone()
            self.assertEqual(progress["mastery_level"], 0)
            self.assertEqual(progress["error_count"], 1)

    def test_knowledge_page_has_delete_dialog_and_rejects_sibling_duplicate_names(self):
        page = self.client.get("/knowledge")
        self.assertEqual(page.status_code, 200)
        self.assertIn(b"data-knowledge-delete-dialog", page.data)
        with self.mod.db() as conn:
            subject = conn.execute("SELECT id FROM subjects WHERE project_id=? LIMIT 1", (self.project_id,)).fetchone()["id"]
            chapter = conn.execute("SELECT id FROM chapters WHERE subject_id=? LIMIT 1", (subject,)).fetchone()["id"]
        self.client.post("/knowledge", data={"kind": "chapter", "subject_id": subject, "name": "测试章节"})
        self.client.post("/knowledge", data={"kind": "point", "chapter_id": chapter, "name": "测试知识点"})
        with self.mod.db() as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) n FROM chapters WHERE subject_id=?", (subject,)).fetchone()["n"], 1)
            self.assertEqual(conn.execute("SELECT COUNT(*) n FROM knowledge_points WHERE chapter_id=?", (chapter,)).fetchone()["n"], 1)
            first = conn.execute("INSERT INTO chapters(subject_id,name,is_core) VALUES (?,?,0)", (subject, "Network")).lastrowid
            second = conn.execute("INSERT INTO chapters(subject_id,name,is_core) VALUES (?,?,0)", (subject, "Routing")).lastrowid
        self.client.post("/knowledge/rename", data={"kind": "chapter", "id": second, "name": "network"})
        with self.mod.db() as conn:
            self.assertEqual(conn.execute("SELECT name FROM chapters WHERE id=?", (second,)).fetchone()["name"], "Routing")
            self.assertEqual(conn.execute("SELECT name FROM chapters WHERE id=?", (first,)).fetchone()["name"], "Network")

    def test_knowledge_tree_uses_browse_first_hierarchy_and_correct_counts(self):
        question_id = self.add_question()
        with self.mod.db() as conn:
            conn.execute("UPDATE questions SET stem=? WHERE id=?", ("层级统计测试题", question_id))
        page = self.client.get("/knowledge")
        self.assertEqual(page.status_code, 200)
        self.assertIn(b"data-knowledge-tree", page.data)
        self.assertIn(b"data-knowledge-node-dialog", page.data)
        self.assertIn("知识结构".encode(), page.data)
        self.assertIn("1 个科目 · 1 个章节 · 1 个知识点 · 1 道题".encode(), page.data)
        self.assertNotIn(b"knowledge-editor", page.data)
        with self.mod.db() as conn:
            point_id = conn.execute(
                "SELECT knowledge_point_id FROM questions WHERE id=?", (question_id,)
            ).fetchone()["knowledge_point_id"]
        self.assertIn(f"/questions?knowledge_point_id={point_id}".encode(), page.data)

    def test_knowledge_point_question_filter_and_chapter_dialog_update(self):
        first_question = self.add_question()
        with self.mod.db() as conn:
            conn.execute("UPDATE questions SET stem=? WHERE id=?", ("只属于第一个知识点", first_question))
            chapter = conn.execute("SELECT id FROM chapters ORDER BY id LIMIT 1").fetchone()["id"]
            first_point = conn.execute(
                "SELECT knowledge_point_id FROM questions WHERE id=?", (first_question,)
            ).fetchone()["knowledge_point_id"]
            second_point = conn.execute(
                "INSERT INTO knowledge_points(chapter_id,name) VALUES (?,?)", (chapter, "第二知识点")
            ).lastrowid
            conn.execute("""INSERT INTO questions(knowledge_point_id,type,stem,options_json,answer_json,status,created_at,updated_at)
              VALUES (?,?,?,?,?,'verified',?,?)""",
              (second_point, "single", "只属于第二个知识点", '["A","B"]', '["A"]', self.mod.now_iso(), self.mod.now_iso()))
        filtered = self.client.get(f"/questions?knowledge_point_id={first_point}")
        self.assertIn("只属于第一个知识点".encode(), filtered.data)
        self.assertNotIn("只属于第二个知识点".encode(), filtered.data)
        response = self.client.post(
            "/knowledge/rename",
            data={"kind": "chapter", "id": chapter, "name": "更新后的章节", "is_core": "on"},
        )
        self.assertIn(f"focus_kind=chapter&focus_id={chapter}", response.headers["Location"])
        with self.mod.db() as conn:
            updated = conn.execute("SELECT name,is_core FROM chapters WHERE id=?", (chapter,)).fetchone()
            self.assertEqual(updated["name"], "更新后的章节")
            self.assertEqual(updated["is_core"], 1)

    def test_empty_knowledge_node_deletes_without_backup(self):
        subject, _, _ = self.create_tree("空科目")
        before = set(self.mod.BACKUP_DIR.glob("knowledge_delete_*")) if self.mod.BACKUP_DIR.exists() else set()
        reports_dir = self.mod.BACKUP_DIR / "knowledge_operation_reports"
        report_count = len(list(reports_dir.glob("*.json"))) if reports_dir.exists() else 0
        response = self.client.post("/knowledge/delete", data={"kind": "subject", "node_id": subject})
        self.assertEqual(response.status_code, 302)
        with self.mod.db() as conn:
            self.assertIsNone(conn.execute("SELECT 1 FROM subjects WHERE id=?", (subject,)).fetchone())
        after = set(self.mod.BACKUP_DIR.glob("knowledge_delete_*")) if self.mod.BACKUP_DIR.exists() else set()
        self.assertEqual(before, after)
        reports = list(reports_dir.glob("*.json"))
        self.assertEqual(len(reports), report_count + 1)
        self.assertLess(max(reports, key=lambda path: path.stat().st_mtime).stat().st_size, 2048)

    def test_nonempty_point_migrates_question_identity_and_creates_lightweight_backup(self):
        with self.mod.db() as conn:
            chapter = conn.execute("SELECT id FROM chapters ORDER BY id LIMIT 1").fetchone()["id"]
            source = conn.execute("INSERT INTO knowledge_points(chapter_id,name) VALUES (?,?)", (chapter, "源知识点")).lastrowid
            target = conn.execute("INSERT INTO knowledge_points(chapter_id,name) VALUES (?,?)", (chapter, "目标知识点")).lastrowid
            question = conn.execute("""INSERT INTO questions(knowledge_point_id,type,stem,options_json,answer_json,status,created_at,updated_at)
              VALUES (?,'single','迁移保留ID','[\"甲\",\"乙\"]','[\"A\"]','verified',?,?)""",
              (source, self.mod.now_iso(), self.mod.now_iso())).lastrowid
            conn.execute("INSERT INTO question_progress(question_id,mastery_level,attempts,correct_attempts,error_count) VALUES (?,?,?,?,?)",
                         (question, 3, 5, 4, 1))
            conn.execute("INSERT INTO attempts(question_id,mode,is_correct,answered_at) VALUES (?,?,?,?)",
                         (question, "practice", 1, self.mod.now_iso()))
            conn.execute("INSERT INTO practice_sessions(id,project_id,mode,question_ids_json,started_at) VALUES (?,?,?,?,?)",
                         ("migration-session", self.project_id, "practice", json.dumps([question]), self.mod.now_iso()))
        impact = self.client.get(f"/api/knowledge/point/{source}/delete-impact").get_json()
        self.assertTrue(impact["nonempty"])
        self.assertEqual(impact["counts"]["questions"], 1)
        response = self.client.post("/knowledge/delete", data={
            "kind": "point", "node_id": source, "target_id": target, "confirmation_name": "源知识点"
        })
        self.assertEqual(response.status_code, 302)
        with self.mod.db() as conn:
            self.assertEqual(conn.execute("SELECT knowledge_point_id FROM questions WHERE id=?", (question,)).fetchone()["knowledge_point_id"], target)
            self.assertEqual(conn.execute("SELECT mastery_level FROM question_progress WHERE question_id=?", (question,)).fetchone()["mastery_level"], 3)
            self.assertEqual(conn.execute("SELECT COUNT(*) n FROM attempts WHERE question_id=?", (question,)).fetchone()["n"], 1)
            self.assertEqual(json.loads(conn.execute("SELECT question_ids_json FROM practice_sessions WHERE id='migration-session'").fetchone()["question_ids_json"]), [question])
        folders = sorted(self.mod.BACKUP_DIR.glob(f"knowledge_delete_*_point_{source}"))
        self.assertTrue(folders)
        summary = folders[-1] / "summary.json"
        self.assertTrue((folders[-1] / self.mod.DB_PATH.name).exists())
        self.assertLess(summary.stat().st_size, 2048)
        self.assertEqual(json.loads(summary.read_text(encoding="utf-8"))["result"], "success")

    def test_chapter_delete_requires_conflict_confirmation_then_merges_and_preserves_core(self):
        with self.mod.db() as conn:
            subject = conn.execute("SELECT id FROM subjects ORDER BY id LIMIT 1").fetchone()["id"]
            source_chapter = conn.execute("INSERT INTO chapters(subject_id,name,is_core) VALUES (?,?,1)", (subject, "源章节")).lastrowid
            target_chapter = conn.execute("INSERT INTO chapters(subject_id,name,is_core) VALUES (?,?,0)", (subject, "目标章节")).lastrowid
            source_point = conn.execute("INSERT INTO knowledge_points(chapter_id,name) VALUES (?,?)", (source_chapter, "同名点")).lastrowid
            target_point = conn.execute("INSERT INTO knowledge_points(chapter_id,name) VALUES (?,?)", (target_chapter, "同名点")).lastrowid
            question = conn.execute("""INSERT INTO questions(knowledge_point_id,type,stem,options_json,answer_json,status,created_at,updated_at)
              VALUES (?,'single','章节迁移题','[\"甲\",\"乙\"]','[\"A\"]','verified',?,?)""",
              (source_point, self.mod.now_iso(), self.mod.now_iso())).lastrowid
            lab = conn.execute("INSERT INTO labs(project_id,chapter_id,title,status,created_at,updated_at) VALUES (?,?,?,?,?,?)",
                               (self.project_id, source_chapter, "迁移实践", "planned", self.mod.now_iso(), self.mod.now_iso())).lastrowid
        blocked = self.client.post("/knowledge/delete", data={"kind": "chapter", "node_id": source_chapter,
            "target_id": target_chapter, "confirmation_name": "源章节"})
        self.assertEqual(blocked.status_code, 302)
        with self.mod.db() as conn:
            self.assertIsNotNone(conn.execute("SELECT 1 FROM chapters WHERE id=?", (source_chapter,)).fetchone())
        self.client.post("/knowledge/delete", data={"kind": "chapter", "node_id": source_chapter,
            "target_id": target_chapter, "confirmation_name": "源章节", "merge_conflicts": "1"})
        with self.mod.db() as conn:
            self.assertIsNone(conn.execute("SELECT 1 FROM chapters WHERE id=?", (source_chapter,)).fetchone())
            self.assertEqual(conn.execute("SELECT knowledge_point_id FROM questions WHERE id=?", (question,)).fetchone()["knowledge_point_id"], target_point)
            self.assertEqual(conn.execute("SELECT chapter_id FROM labs WHERE id=?", (lab,)).fetchone()["chapter_id"], target_chapter)
            self.assertEqual(conn.execute("SELECT is_core FROM chapters WHERE id=?", (target_chapter,)).fetchone()["is_core"], 1)

    def test_existing_duplicate_children_block_migration_until_resolved(self):
        with self.mod.db() as conn:
            subject = conn.execute("SELECT id FROM subjects ORDER BY id LIMIT 1").fetchone()["id"]
            source = conn.execute("INSERT INTO chapters(subject_id,name,is_core) VALUES (?,?,0)", (subject, "含重名章节")).lastrowid
            target = conn.execute("INSERT INTO chapters(subject_id,name,is_core) VALUES (?,?,0)", (subject, "空目标章节")).lastrowid
            conn.execute("INSERT INTO knowledge_points(chapter_id,name) VALUES (?,?)", (source, "Duplicate"))
            conn.execute("INSERT INTO knowledge_points(chapter_id,name) VALUES (?,?)", (source, "duplicate"))
        impact = self.client.get(f"/api/knowledge/chapter/{source}/delete-impact?target_id={target}").get_json()
        self.assertTrue(impact["has_ambiguous_conflicts"])
        self.client.post("/knowledge/delete", data={"kind": "chapter", "node_id": source,
            "target_id": target, "confirmation_name": "含重名章节", "merge_conflicts": "1"})
        with self.mod.db() as conn:
            self.assertIsNotNone(conn.execute("SELECT 1 FROM chapters WHERE id=?", (source,)).fetchone())

    def test_subject_delete_blocks_active_import_then_moves_import_history(self):
        source_subject, source_chapter, _ = self.create_tree("待删除科目", "待迁移章节", "待迁移知识点")
        target_subject, _, _ = self.create_tree("目标科目")
        with self.mod.db() as conn:
            document = "delete-source-doc"
            job = "delete-source-job"
            conn.execute("""INSERT INTO source_documents(id,project_id,subject_id,original_name,stored_path,sha256,file_type,size_bytes,created_at)
              VALUES (?,?,?,?,?,?,?,?,?)""", (document, self.project_id, source_subject, "source.csv", "imports/source.csv", "abc", "csv", 1, self.mod.now_iso()))
            conn.execute("""INSERT INTO import_jobs(id,source_document_id,project_id,subject_id,status,stage,created_at,updated_at)
              VALUES (?,?,?,?,?,?,?,?)""", (job, document, self.project_id, source_subject, "running", "parsing", self.mod.now_iso(), self.mod.now_iso()))
        self.client.post("/knowledge/delete", data={"kind": "subject", "node_id": source_subject,
            "target_id": target_subject, "confirmation_name": "待删除科目"})
        with self.mod.db() as conn:
            self.assertIsNotNone(conn.execute("SELECT 1 FROM subjects WHERE id=?", (source_subject,)).fetchone())
            conn.execute("UPDATE import_jobs SET status='failed' WHERE id=?", (job,))
        self.client.post("/knowledge/delete", data={"kind": "subject", "node_id": source_subject,
            "target_id": target_subject, "confirmation_name": "待删除科目"})
        with self.mod.db() as conn:
            self.assertIsNone(conn.execute("SELECT 1 FROM subjects WHERE id=?", (source_subject,)).fetchone())
            self.assertEqual(conn.execute("SELECT subject_id FROM chapters WHERE id=?", (source_chapter,)).fetchone()["subject_id"], target_subject)
            self.assertEqual(conn.execute("SELECT subject_id FROM source_documents WHERE id=?", (document,)).fetchone()["subject_id"], target_subject)
            self.assertEqual(conn.execute("SELECT subject_id FROM import_jobs WHERE id=?", (job,)).fetchone()["subject_id"], target_subject)

    def test_delete_target_cannot_cross_project(self):
        with self.mod.db() as conn:
            source = conn.execute("SELECT id FROM knowledge_points ORDER BY id LIMIT 1").fetchone()["id"]
            question = conn.execute("""INSERT INTO questions(knowledge_point_id,type,stem,options_json,answer_json,status,created_at,updated_at)
              VALUES (?,'single','跨项目保护','[\"甲\",\"乙\"]','[\"A\"]','verified',?,?)""",
              (source, self.mod.now_iso(), self.mod.now_iso())).lastrowid
            other = self.mod.create_project(conn, "另一个项目", "practice")
            subject = conn.execute("INSERT INTO subjects(project_id,name,code) VALUES (?,?,?)", (other, "其它科目", "")).lastrowid
            chapter = conn.execute("INSERT INTO chapters(subject_id,name,is_core) VALUES (?,?,0)", (subject, "其它章节")).lastrowid
            target = conn.execute("INSERT INTO knowledge_points(chapter_id,name) VALUES (?,?)", (chapter, "其它知识点")).lastrowid
        response = self.client.post("/knowledge/delete", data={"kind": "point", "node_id": source,
            "target_id": target, "confirmation_name": "测试知识点"})
        self.assertEqual(response.status_code, 302)
        with self.mod.db() as conn:
            self.assertEqual(conn.execute("SELECT knowledge_point_id FROM questions WHERE id=?", (question,)).fetchone()["knowledge_point_id"], source)

    def test_last_empty_subject_can_be_deleted_and_page_stays_usable(self):
        with self.mod.db() as conn:
            project = self.mod.create_project(conn, "空项目", "practice")
            subject = conn.execute("INSERT INTO subjects(project_id,name,code) VALUES (?,?,?)", (project, "最后科目", "")).lastrowid
        self.client.post("/projects/switch", data={"project_id": project, "return_to": "/knowledge"})
        self.client.post("/knowledge/delete", data={"kind": "subject", "node_id": subject})
        page = self.client.get("/knowledge")
        self.assertEqual(page.status_code, 200)
        self.assertIn("请先新增科目".encode(), page.data)

    def test_knowledge_backup_retention_keeps_latest_twenty(self):
        from knowledge_service import _trim_backups
        folder = Path(self.temp.name) / "retention-test"
        folder.mkdir(exist_ok=True)
        for index in range(22):
            item = folder / f"knowledge_delete_{index:02d}"
            item.mkdir(exist_ok=True)
            os.utime(item, (index + 1, index + 1))
        _trim_backups(folder, keep=20)
        remaining = sorted(path.name for path in folder.glob("knowledge_delete_*"))
        self.assertEqual(len(remaining), 20)
        self.assertNotIn("knowledge_delete_00", remaining)
        self.assertNotIn("knowledge_delete_01", remaining)

    def test_subjective_question_requires_self_rating(self):
        qid = self.add_question("short")
        with self.mod.db() as conn:
            sid = "subjective-test"
            conn.execute("INSERT INTO practice_sessions(id,project_id,mode,question_ids_json,started_at) VALUES (?,?,?,?,?)", (sid, self.project_id, "practice", json.dumps([qid]), self.mod.now_iso()))
        self.client.get(f"/practice/{sid}")
        reveal = self.client.post(f"/practice/{sid}", data={"action": "reveal", "control_token": self.control_token("practice_sessions", sid)})
        self.assertIn("请对照参考答案自评".encode("utf-8"), reveal.data)
        rated = self.client.post(f"/practice/{sid}", data={"rating": "fuzzy", "control_token": self.control_token("practice_sessions", sid)})
        self.assertIn("本题已记录".encode("utf-8"), rated.data)
        with self.mod.db() as conn:
            attempt = conn.execute("SELECT * FROM attempts WHERE question_id=?", (qid,)).fetchone()
            self.assertEqual(attempt["self_rating"], "fuzzy")

    def test_subjective_self_rating_stage_survives_pause(self):
        qid = self.add_question("short")
        with self.mod.db() as conn:
            sid = "subjective-pause"
            conn.execute("INSERT INTO practice_sessions(id,project_id,mode,question_ids_json,started_at) VALUES (?,?,?,?,?)",
                         (sid, self.project_id, "practice", json.dumps([qid]), self.mod.now_iso()))
        location = f"/practice/{sid}"
        self.client.get(location)
        token = self.control_token("practice_sessions", sid)
        self.client.post(location, data={"action": "reveal", "draft": "我的草稿", "control_token": token})
        paused = self.client.post(f"/practice/{sid}/pause", json={"control_token": token})
        self.assertEqual(paused.status_code, 200)
        restored = self.client.get(location)
        self.assertIn("请对照参考答案自评".encode(), restored.data)
        self.assertIn("练习已暂停".encode(), restored.data)
        resumed_token = self.control_token("practice_sessions", sid)
        self.client.post(f"/practice/{sid}/resume", data={"control_token": resumed_token})
        self.client.get(location)
        rated = self.client.post(location, data={"rating": "fuzzy", "control_token": self.control_token("practice_sessions", sid)})
        self.assertIn("本题已记录".encode(), rated.data)

    def test_active_session_is_visible_on_dashboard_and_hides_new_practice_form(self):
        self.add_question("single", ["A"])
        location = self.client.post("/practice", data={"count": 1}).headers["Location"]
        dashboard = self.client.get("/")
        self.assertIn("进行中任务".encode(), dashboard.data)
        self.assertIn(location.encode(), dashboard.data)
        setup = self.client.get("/practice")
        self.assertIn("尚未结束的章节练习".encode(), setup.data)
        self.assertNotIn(b'id="practice-setup-form"', setup.data)

    def test_mock_exam_is_snapshotted_and_scored(self):
        qid = self.add_question("single", ["A"])
        response = self.client.post("/mock", data={"count": 1, "minutes": 5})
        self.assertEqual(response.status_code, 302)
        location = response.headers["Location"]
        page = self.client.get(location)
        self.assertIn("模拟考试".encode("utf-8"), page.data)
        exam_id = location.rstrip("/").split("/")[-1]
        result = self.client.post(location, data={f"q_{qid}": "A", "control_token": self.control_token("mock_exams", exam_id)}, follow_redirects=True)
        self.assertIn(b"100.0%", result.data)
        with self.mod.db() as conn:
            exam = conn.execute("SELECT * FROM mock_exams").fetchone()
            self.assertEqual(exam["score"], 100.0)
            self.assertEqual(json.loads(exam["question_ids_json"]), [qid])

    def test_mock_autosave_pause_resume_and_submit_preserve_qualifying(self):
        qid = self.add_question("single", ["A"])
        with self.mod.db() as conn:
            conn.execute("""INSERT INTO project_settings(project_id,key,value) VALUES (?, 'qualifying_count','1')
              ON CONFLICT(project_id,key) DO UPDATE SET value='1'""", (self.project_id,))
            conn.execute("""INSERT INTO project_settings(project_id,key,value) VALUES (?, 'qualifying_minutes','5')
              ON CONFLICT(project_id,key) DO UPDATE SET value='5'""", (self.project_id,))
        location = self.client.post("/mock", data={"count": 1, "minutes": 5}).headers["Location"]
        exam_id = location.rstrip("/").split("/")[-1]
        self.client.get(location)
        token = self.control_token("mock_exams", exam_id)
        saved = self.client.post(f"/api/mock/{exam_id}/answers", json={
            "control_token": token, "question_id": qid, "answers": ["A"],
        })
        self.assertEqual(saved.get_json()["status"], "saved")
        with self.mod.db() as conn:
            conn.execute("UPDATE mock_exams SET active_started_at=? WHERE id=?",
                         ((self.mod.datetime.now() - self.mod.timedelta(seconds=10)).replace(microsecond=0).isoformat(sep=" "), exam_id))
        paused = self.client.post(f"/mock/{exam_id}/pause", data={"control_token": token, f"q_{qid}": "A"})
        self.assertEqual(paused.status_code, 302)
        with self.mod.db() as conn:
            exam = conn.execute("SELECT * FROM mock_exams WHERE id=?", (exam_id,)).fetchone()
            self.assertEqual(exam["status"], "paused")
            self.assertLessEqual(exam["remaining_seconds"], 290)
            qualifying = exam["qualifying"]
            self.assertEqual(qualifying, 1)
        page = self.client.get(location)
        self.assertIn("模拟考试已暂停".encode(), page.data)
        self.assertIn(b'value="A" data-question-id', page.data)
        self.assertNotIn(b"interval=setInterval", page.data)
        self.assertNotIn(b"heartbeat:true", page.data)
        resumed = self.client.post(f"/mock/{exam_id}/resume", data={
            "control_token": self.control_token("mock_exams", exam_id),
        })
        self.assertEqual(resumed.status_code, 302)
        self.client.get(location)
        token = self.control_token("mock_exams", exam_id)
        result = self.client.post(location, data={"control_token": token, f"q_{qid}": "A"}, follow_redirects=True)
        self.assertIn(b"100.0%", result.data)
        with self.mod.db() as conn:
            exam = conn.execute("SELECT * FROM mock_exams WHERE id=?", (exam_id,)).fetchone()
            self.assertEqual(exam["qualifying"], qualifying)
            self.assertEqual(exam["status"], "submitted")

    def test_mock_unpaused_timeout_submits_saved_answers(self):
        qid = self.add_question("single", ["A"])
        location = self.client.post("/mock", data={"count": 1, "minutes": 1}).headers["Location"]
        exam_id = location.rstrip("/").split("/")[-1]
        self.client.get(location)
        token = self.control_token("mock_exams", exam_id)
        self.client.post(f"/api/mock/{exam_id}/answers", json={
            "control_token": token, "question_id": qid, "answers": ["A"],
        })
        with self.mod.db() as conn:
            conn.execute("UPDATE mock_exams SET remaining_seconds=1,active_started_at=? WHERE id=?",
                         ((self.mod.datetime.now() - self.mod.timedelta(seconds=5)).replace(microsecond=0).isoformat(sep=" "), exam_id))
        result = self.client.get(location, follow_redirects=True)
        self.assertIn(b"100.0%", result.data)
        with self.mod.db() as conn:
            exam = conn.execute("SELECT * FROM mock_exams WHERE id=?", (exam_id,)).fetchone()
            self.assertEqual(exam["status"], "submitted")
            self.assertEqual(conn.execute("SELECT COUNT(*) n FROM attempts WHERE session_id=?", (exam_id,)).fetchone()["n"], 1)

    def test_mock_terminate_discards_draft_without_result_or_attempts(self):
        qid = self.add_question("single", ["A"])
        location = self.client.post("/mock", data={"count": 1, "minutes": 5}).headers["Location"]
        exam_id = location.rstrip("/").split("/")[-1]
        self.client.get(location)
        token = self.control_token("mock_exams", exam_id)
        self.client.post(f"/api/mock/{exam_id}/answers", json={
            "control_token": token, "question_id": qid, "answers": ["A"],
        })
        result = self.client.post(f"/mock/{exam_id}/terminate", data={"control_token": token}, follow_redirects=True)
        self.assertIn("不生成成绩".encode(), result.data)
        with self.mod.db() as conn:
            exam = conn.execute("SELECT * FROM mock_exams WHERE id=?", (exam_id,)).fetchone()
            self.assertEqual(exam["status"], "terminated")
            self.assertEqual(exam["answers_json"], "{}")
            self.assertIsNone(exam["submitted_at"])
            self.assertEqual(conn.execute("SELECT COUNT(*) n FROM attempts WHERE session_id=?", (exam_id,)).fetchone()["n"], 0)

    def test_second_device_takes_over_mock_control(self):
        qid = self.add_question("single", ["A"])
        location = self.client.post("/mock", data={"count": 1, "minutes": 5}).headers["Location"]
        exam_id = location.rstrip("/").split("/")[-1]
        first = self.mod.app.test_client(); second = self.mod.app.test_client()
        with first.session_transaction() as state: state["current_project_id"] = self.project_id
        with second.session_transaction() as state: state["current_project_id"] = self.project_id
        first.get(location); first_token = self.control_token("mock_exams", exam_id)
        second.get(location); second_token = self.control_token("mock_exams", exam_id)
        stale = first.post(f"/api/mock/{exam_id}/answers", json={
            "control_token": first_token, "question_id": qid, "answers": ["A"],
        })
        self.assertEqual(stale.status_code, 409)
        self.assertEqual(stale.get_json()["code"], "session_taken_over")
        current = second.post(f"/api/mock/{exam_id}/answers", json={
            "control_token": second_token, "question_id": qid, "answers": ["A"],
        })
        self.assertEqual(current.status_code, 200)

    def test_readiness_requires_all_three_gates(self):
        qid = self.add_question("single", ["A"])
        with self.mod.db() as conn:
            conn.execute("INSERT INTO question_progress(question_id,mastery_level,due_date,attempts,correct_attempts,error_count) VALUES (?,?,?,?,?,?)", (qid, 4, "2099-01-01", 4, 4, 0))
            chapter = conn.execute("SELECT c.id FROM chapters c JOIN knowledge_points kp ON kp.chapter_id=c.id JOIN questions q ON q.knowledge_point_id=kp.id WHERE q.id=?", (qid,)).fetchone()["id"]
            conn.execute("INSERT INTO labs(project_id,chapter_id,title,status,created_at,updated_at) VALUES (?,?,?,?,?,?)", (self.project_id, chapter, "达标实验", "completed", self.mod.now_iso(), self.mod.now_iso()))
            for n in range(3):
                conn.execute("""INSERT INTO mock_exams(id,project_id,started_at,submitted_at,question_ids_json,
                  score,objective_count,time_limit,qualifying,status) VALUES (?,?,?,?,?,?,?,?,1,'submitted')""",
                  (f"ready-{n}", self.project_id, self.mod.now_iso(), self.mod.now_iso(), "[]", 90, 50, 60))
            state = self.mod.readiness(conn, self.project_id)
            self.assertTrue(state["ready"])

    def test_two_browser_clients_keep_independent_current_projects(self):
        with self.mod.db() as conn:
            other = self.mod.create_project(conn, "英语课程", "practice")
        first_client = self.mod.app.test_client()
        second_client = self.mod.app.test_client()
        first_client.post("/projects/switch", data={"project_id": self.project_id, "return_to": "/"})
        second_client.post("/projects/switch", data={"project_id": other, "return_to": "/"})
        self.assertIn("HCL实验".encode(), first_client.get("/").data)
        self.assertIn("英语课程".encode(), second_client.get("/").data)
        with first_client.session_transaction() as first_session, second_client.session_transaction() as second_session:
            self.assertEqual(first_session["current_project_id"], self.project_id)
            self.assertEqual(second_session["current_project_id"], other)

    def test_question_lists_and_direct_session_urls_are_project_isolated(self):
        first_question = self.add_question()
        with self.mod.db() as conn:
            other = self.mod.create_project(conn, "软考", "exam_prep")
            subject = conn.execute("INSERT INTO subjects(project_id,name,code) VALUES (?,?,?)", (other, "网络工程师", "RK-NET")).lastrowid
            chapter = conn.execute("INSERT INTO chapters(subject_id,name,is_core) VALUES (?,?,1)", (subject, "网络基础")).lastrowid
            point = conn.execute("INSERT INTO knowledge_points(chapter_id,name) VALUES (?,?)", (chapter, "OSI")).lastrowid
            other_question = conn.execute("""INSERT INTO questions(knowledge_point_id,type,stem,options_json,answer_json,
              status,created_at,updated_at) VALUES (?,?,?,?,?,'verified',?,?)""",
              (point, "single", "只属于软考项目", '["A","B"]', '["A"]', self.mod.now_iso(), self.mod.now_iso())).lastrowid
            sid = "other-project-session"
            conn.execute("INSERT INTO practice_sessions(id,project_id,mode,question_ids_json,started_at) VALUES (?,?,?,?,?)",
                         (sid, other, "practice", json.dumps([other_question]), self.mod.now_iso()))
        page = self.client.get("/questions")
        self.assertNotIn("只属于软考项目".encode(), page.data)
        self.assertEqual(self.client.get(f"/questions/{other_question}/edit").status_code, 404)
        self.assertEqual(self.client.get(f"/practice/{sid}").status_code, 404)
        self.assertEqual(self.client.get(f"/questions/{first_question}/edit").status_code, 200)

    def test_csv_import_uses_confirm_parse_commit_pipeline_and_creates_draft(self):
        job_id, job = self.upload_csv_job()
        self.assertEqual(job["status"], "ready")
        self.assertEqual(job["candidate_count"], 1)
        result = self.client.post(f"/imports/{job_id}/commit", data={"duplicate_action": "skip"})
        self.assertEqual(result.status_code, 302)
        with self.mod.db() as conn:
            question = conn.execute("SELECT * FROM questions WHERE stem='通用导入测试题'").fetchone()
            self.assertIsNotNone(question)
            self.assertEqual(question["status"], "draft")
            self.assertEqual(question["source_item_key"], "CSV-1")
            self.assertEqual(question["import_batch_id"], job_id)

    def test_duplicate_import_defaults_to_skip_and_can_keep_copy(self):
        first_id, first = self.upload_csv_job("重复题测试")
        self.assertEqual(first["status"], "ready")
        self.client.post(f"/imports/{first_id}/commit", data={"duplicate_action": "skip"})
        second_id, second = self.upload_csv_job("重复题测试")
        self.assertEqual(second["duplicate_count"], 1)
        self.client.post(f"/imports/{second_id}/commit", data={"duplicate_action": "skip"})
        with self.mod.db() as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) n FROM questions WHERE stem='重复题测试'").fetchone()["n"], 1)
        third_id, third = self.upload_csv_job("重复题测试")
        self.assertEqual(third["duplicate_count"], 1)
        self.client.post(f"/imports/{third_id}/commit", data={"duplicate_action": "copy"})
        with self.mod.db() as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) n FROM questions WHERE stem='重复题测试'").fetchone()["n"], 2)
        fourth_id, fourth = self.upload_csv_job("重复题测试")
        self.assertGreaterEqual(fourth["duplicate_count"], 1)
        self.client.post(f"/imports/{fourth_id}/commit", data={"duplicate_action": "update"})
        with self.mod.db() as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) n FROM questions WHERE stem='重复题测试'").fetchone()["n"], 2)
            self.assertIsNotNone(conn.execute("SELECT 1 FROM questions WHERE stem='重复题测试' AND import_batch_id=?", (fourth_id,)).fetchone())

    def test_docx_and_xlsx_fixed_formats_parse(self):
        from docx import Document
        from openpyxl import Workbook
        from import_service import parse_document

        root = Path(self.temp.name)
        docx_path = root / "sample.docx"
        document = Document()
        for line in ("1. DOCX 测试题", "A. 正确", "B. 错误", "答案：A", "解析：DOCX 解析"):
            document.add_paragraph(line)
        document.save(docx_path)
        docx_records, docx_errors = parse_document(docx_path, "docx")
        self.assertEqual(len(docx_records), 1)
        self.assertFalse(docx_records[0]["validation_error"])
        self.assertFalse(docx_errors)

        xlsx_path = root / "sample.xlsx"
        workbook = Workbook(); sheet = workbook.active
        sheet.append(["题型", "题干", "选项A", "选项B", "答案", "解析", "章节", "知识点", "原题号"])
        sheet.append(["单选", "XLSX 测试题", "正确", "错误", "A", "解析", "章节", "知识点", "X-1"])
        workbook.save(xlsx_path); workbook.close()
        xlsx_records, xlsx_errors = parse_document(xlsx_path, "xlsx")
        self.assertEqual(len(xlsx_records), 1)
        self.assertFalse(xlsx_records[0]["validation_error"])
        self.assertFalse(xlsx_errors)

    def test_unknown_import_format_is_rejected_without_questions(self):
        response = self.client.post("/imports", data={"file": (io.BytesIO(b"legacy"), "legacy.doc")}, content_type="multipart/form-data")
        self.assertEqual(response.status_code, 302)
        job_id = response.headers["Location"].rstrip("/").split("/")[-1]
        with self.mod.db() as conn:
            job = conn.execute("SELECT * FROM import_jobs WHERE id=?", (job_id,)).fetchone()
            self.assertEqual(job["status"], "failed")
            self.assertEqual(conn.execute("SELECT COUNT(*) n FROM questions").fetchone()["n"], 0)

    def test_invalid_import_batch_never_partially_writes(self):
        csv_data = "题型,题干,选项A,选项B,答案,解析,章节,知识点,原题号\n单选,有效题,甲,乙,A,解析,章节,知识点,1\n单选,缺答案题,甲,乙,,解析,章节,知识点,2\n"
        response = self.client.post("/imports", data={"file": (io.BytesIO(csv_data.encode("utf-8-sig")), "invalid.csv")}, content_type="multipart/form-data")
        job_id = response.headers["Location"].rstrip("/").split("/")[-1]
        self.wait_for_job(job_id, {"waiting_target"})
        with self.mod.db() as conn:
            subject_id = conn.execute("SELECT id FROM subjects WHERE project_id=? ORDER BY id LIMIT 1", (self.project_id,)).fetchone()["id"]
        self.client.post(f"/imports/{job_id}/target", data={"project_id": self.project_id, "subject_id": subject_id})
        job = self.wait_for_job(job_id, {"blocked", "ready", "failed"})
        self.assertEqual(job["status"], "blocked")
        self.client.post(f"/imports/{job_id}/commit", data={"duplicate_action": "skip"})
        with self.mod.db() as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) n FROM questions").fetchone()["n"], 0)

    def test_export_scope_status_filters_and_count_api(self):
        self.add_question(status="verified")
        self.add_question(qtype="multiple", status="draft")
        self.add_question(qtype="short", status="archived")
        with self.mod.db() as conn:
            subject_id = conn.execute("SELECT id FROM subjects WHERE project_id=?", (self.project_id,)).fetchone()["id"]
            chapter_id = conn.execute("SELECT id FROM chapters WHERE subject_id=?", (subject_id,)).fetchone()["id"]
        self.assertEqual(self.client.get("/api/exports/count?scope_type=project").get_json()["count"], 1)
        self.assertEqual(self.client.get("/api/exports/count?scope_type=project&include_drafts=1").get_json()["count"], 2)
        self.assertEqual(self.client.get(f"/api/exports/count?scope_type=subject&scope_id={subject_id}&include_drafts=1").get_json()["count"], 2)
        self.assertEqual(self.client.get(f"/api/exports/count?scope_type=chapter&scope_id={chapter_id}&include_drafts=1").get_json()["count"], 2)
        job_id, job = self.create_export("chapter", chapter_id, include_drafts=True)
        self.assertEqual(job["status"], "completed", job["message"])
        self.assertEqual(job["question_count"], 2)
        download = self.client.get(f"/exports/{job_id}/download")
        self.assertEqual(download.status_code, 200)
        self.assertEqual(download.mimetype, "application/zip")
        download.close()

    def test_failed_export_can_retry_then_delete_without_touching_questions(self):
        self.add_question(status="verified")
        original_limit = self.mod.export_queue.max_bytes
        try:
            self.mod.export_queue.max_bytes = 1
            job_id, job = self.create_export()
            self.assertEqual(job["status"], "failed")
            self.assertIn("上限", job["message"])
            self.mod.export_queue.max_bytes = original_limit
            self.client.post(f"/exports/{job_id}/retry")
            completed = self.wait_for_export(job_id, {"completed", "failed"})
            self.assertEqual(completed["status"], "completed", completed["message"])
            package_path = Path(self.temp.name) / completed["stored_path"]
            self.assertTrue(package_path.is_file())
            self.client.post(f"/exports/{job_id}/delete")
            self.assertFalse(package_path.exists())
            with self.mod.db() as conn:
                self.assertIsNone(conn.execute("SELECT 1 FROM export_jobs WHERE id=?", (job_id,)).fetchone())
                self.assertEqual(conn.execute("SELECT COUNT(*) n FROM questions").fetchone()["n"], 1)
        finally:
            self.mod.export_queue.max_bytes = original_limit

    def test_share_package_round_trip_preserves_content_but_not_personal_data(self):
        ids = [self.add_question(qtype=qtype, status="draft" if qtype == "short" else "verified")
               for qtype in ("single", "multiple", "true_false", "fill", "short")]
        self.add_question(qtype="single", status="archived")
        image_name = "roundtrip.png"
        (Path(self.temp.name) / "uploads").mkdir(exist_ok=True)
        (Path(self.temp.name) / "uploads" / image_name).write_bytes(b"test-image-content")
        with self.mod.db() as conn:
            conn.execute("UPDATE questions SET image_path=? WHERE id=?", (image_name, ids[0]))
            conn.execute("INSERT INTO question_progress(question_id,mastery_level,attempts,correct_attempts,error_count) VALUES (?,?,?,?,?)",
                         (ids[0], 3, 5, 4, 1))
            conn.execute("INSERT INTO attempts(question_id,mode,is_correct,answered_at) VALUES (?,?,?,?)",
                         (ids[0], "practice", 1, self.mod.now_iso()))
        export_id, export_job = self.create_export(include_drafts=True)
        self.assertEqual(export_job["status"], "completed", export_job["message"])
        package_path = Path(self.temp.name) / export_job["stored_path"]
        from transfer_service import inspect_share_package
        manifest, records = inspect_share_package(package_path)
        self.assertEqual(len(records), 5)
        self.assertEqual({item["type"] for item in records}, {"single", "multiple", "true_false", "fill", "short"})
        self.assertEqual({item["status"] for item in records}, {"verified", "draft"})
        self.assertNotIn("id", records[0])
        for forbidden in ("source", "version_note", "source_page", "source_document_id", "import_batch_id", "parser_version"):
            self.assertNotIn(forbidden, records[0])
        self.assertEqual(manifest["counts"]["images"], 1)

        response = self.client.post("/imports", data={"file": (io.BytesIO(package_path.read_bytes()), "shared.zip")},
                                    content_type="multipart/form-data")
        import_id = response.headers["Location"].rstrip("/").split("/")[-1]
        self.wait_for_job(import_id, {"waiting_target", "failed"})
        self.assertEqual(self.client.get(f"/imports/{import_id}").status_code, 200)
        self.client.post(f"/imports/{import_id}/target", data={"new_project_name": "接收项目"})
        mapped = self.wait_for_job(import_id, {"waiting_mapping", "failed"})
        self.assertEqual(mapped["status"], "waiting_mapping", mapped["message"])
        self.assertEqual(self.client.get(f"/imports/{import_id}").status_code, 200)
        detected = json.loads(mapped["detected_json"])
        mapping_data = {f"map_{source['key']}": "new" for source in detected["package"]["subjects"]}
        self.client.post(f"/imports/{import_id}/mapping", data=mapping_data)
        ready = self.wait_for_job(import_id, {"ready", "blocked", "failed"})
        self.assertEqual(ready["status"], "ready", ready["message"])
        self.assertEqual(self.client.get(f"/imports/{import_id}").status_code, 200)
        self.client.post(f"/imports/{import_id}/commit", data={"duplicate_action": "skip", "status_strategy": "preserve"})
        with self.mod.db() as conn:
            target_project = conn.execute("SELECT id FROM learning_projects WHERE name='接收项目'").fetchone()["id"]
            imported = conn.execute("""SELECT q.* FROM questions q JOIN knowledge_points kp ON kp.id=q.knowledge_point_id
              JOIN chapters c ON c.id=kp.chapter_id JOIN subjects s ON s.id=c.subject_id
              WHERE s.project_id=? ORDER BY q.id""", (target_project,)).fetchall()
            self.assertEqual(len(imported), 5)
            self.assertEqual({row["type"] for row in imported}, {"single", "multiple", "true_false", "fill", "short"})
            self.assertEqual({row["status"] for row in imported}, {"verified", "draft"})
            self.assertTrue(all(row["source"] == "" and row["version_note"] == "" for row in imported))
            self.assertEqual(conn.execute("""SELECT COUNT(*) n FROM question_progress qp JOIN questions q ON q.id=qp.question_id
              JOIN knowledge_points kp ON kp.id=q.knowledge_point_id JOIN chapters c ON c.id=kp.chapter_id
              JOIN subjects s ON s.id=c.subject_id WHERE s.project_id=?""", (target_project,)).fetchone()["n"], 0)
            imported_image = next(row["image_path"] for row in imported if row["image_path"])
        self.assertTrue((Path(self.temp.name) / "uploads" / imported_image).is_file())

    def test_package_duplicate_update_preserves_receiver_id_and_progress(self):
        question_id = self.add_question(status="verified")
        with self.mod.db() as conn:
            conn.execute("INSERT INTO question_progress(question_id,mastery_level,attempts,correct_attempts,error_count) VALUES (?,?,?,?,?)",
                         (question_id, 4, 8, 7, 1))
            subject_id = conn.execute("SELECT id FROM subjects WHERE project_id=?", (self.project_id,)).fetchone()["id"]
        _, export_job = self.create_export()
        package_path = Path(self.temp.name) / export_job["stored_path"]
        response = self.client.post("/imports", data={"file": (io.BytesIO(package_path.read_bytes()), "duplicate.zip")},
                                    content_type="multipart/form-data")
        import_id = response.headers["Location"].rstrip("/").split("/")[-1]
        self.wait_for_job(import_id, {"waiting_target"})
        self.client.post(f"/imports/{import_id}/target", data={"project_id": self.project_id})
        mapped = self.wait_for_job(import_id, {"waiting_mapping"})
        source_key = json.loads(mapped["detected_json"])["package"]["subjects"][0]["key"]
        self.client.post(f"/imports/{import_id}/mapping", data={f"map_{source_key}": str(subject_id)})
        ready = self.wait_for_job(import_id, {"ready", "failed"})
        self.assertEqual(ready["duplicate_count"], 1)
        self.client.post(f"/imports/{import_id}/commit", data={"duplicate_action": "update", "status_strategy": "draft"})
        with self.mod.db() as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) n FROM questions").fetchone()["n"], 1)
            self.assertEqual(conn.execute("SELECT id FROM questions").fetchone()["id"], question_id)
            progress = conn.execute("SELECT mastery_level,attempts FROM question_progress WHERE question_id=?", (question_id,)).fetchone()
            self.assertEqual((progress["mastery_level"], progress["attempts"]), (4, 8))
            self.assertEqual(conn.execute("SELECT status FROM questions WHERE id=?", (question_id,)).fetchone()["status"], "draft")

    def test_disabled_optional_module_has_clear_direct_access_page(self):
        with self.mod.db() as conn:
            other = self.mod.create_project(conn, "纯刷题项目", "practice")
        self.client.post("/projects/switch", data={"project_id": other, "return_to": "/"})
        page = self.client.get("/mock")
        self.assertEqual(page.status_code, 200)
        self.assertIn("当前项目未启用".encode(), page.data)
        self.assertNotIn("模拟考试</span>".encode(), self.client.get("/").data)

    def test_archive_restore_and_permanent_delete_create_backup(self):
        with self.mod.db() as conn:
            other = self.mod.create_project(conn, "可归档项目", "practice")
        self.client.post("/projects/switch", data={"project_id": other, "return_to": "/"})
        self.client.post(f"/projects/{other}/archive")
        with self.client.session_transaction() as browser_session:
            self.assertEqual(browser_session["current_project_id"], self.project_id)
        self.client.post(f"/projects/{other}/restore")
        self.client.post(f"/projects/{other}/archive")
        response = self.client.post(f"/projects/{other}/delete", data={"confirmation": "永久删除"})
        self.assertEqual(response.status_code, 302)
        with self.mod.db() as conn:
            self.assertIsNone(conn.execute("SELECT 1 FROM learning_projects WHERE id=?", (other,)).fetchone())
        self.assertTrue(any(self.mod.BACKUP_DIR.glob(f"pre_delete_project_{other}_*/data/{self.mod.DB_PATH.name}")))


if __name__ == "__main__":
    unittest.main()
